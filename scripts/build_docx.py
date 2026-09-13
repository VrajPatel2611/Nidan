#!/usr/bin/env python3
"""
Generate the .docx copies of the Markdown documentation.

    pip install -e ".[docs]"
    python scripts/build_docx.py                 # every document
    python scripts/build_docx.py PRD BUILD_PLAN  # only those

The Markdown is the source of truth. These .docx files exist because they are
what gets shared with people who do not read a repository — a mentor, a
clinical reviewer, an examiner — and they were previously produced by hand,
one at a time, which is why seven build logs had none and most of the rest had
drifted from the Markdown they were made from.

Regenerate after changing any document. The output is deterministic apart from
the generated-on date, so a diff shows only what actually changed.
"""

from __future__ import annotations

import pathlib
import re
import sys
from datetime import date

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError as e:                                  # pragma: no cover
    raise SystemExit(
        'python-docx is not installed. Run:  pip install -e ".[docs]"') from e

REPO = pathlib.Path(__file__).resolve().parent.parent

# Markdown source -> .docx name. Explicit rather than derived: the existing
# names were chosen by hand, and inferring them would rename half the files the
# first time this ran.
DOCUMENTS: dict[str, str] = {
    # specification
    "docs/spec/PRD.md":                 "docs/spec/Nidan_PRD.docx",
    "docs/spec/TECH_SPEC.md":           "docs/spec/Nidan_Tech_Spec.docx",
    "docs/spec/DATA_MODEL.md":          "docs/spec/Nidan_Data_Model.docx",
    "docs/spec/API_CONTRACT.md":        "docs/spec/Nidan_API_Contract.docx",
    "docs/spec/UX_SPEC.md":             "docs/spec/Nidan_UX_Spec.docx",
    "docs/spec/BUILD_PLAN.md":          "docs/spec/Nidan_Build_Plan.docx",
    "docs/spec/TEST_STRATEGY.md":       "docs/spec/Nidan_Test_Strategy.docx",
    "docs/spec/SECURITY_SPEC.md":       "docs/spec/Nidan_Security_Spec.docx",
    # process
    "docs/process/COMMANDS.md":         "docs/process/Nidan_Commands.docx",
    "docs/process/WINDOWS_SETUP.md":    "docs/process/Nidan_Windows_Setup.docx",
    "docs/process/AI_BUILD_PROMPT.md":  "docs/process/Nidan_AI_Build_Prompt.docx",
    "docs/process/USAGE.md":            "docs/process/Nidan_AI_Build_Prompt_Usage.docx",
    "docs/process/RENAME_PLAN.md":      "docs/process/Nidan_Rename_Plan.docx",
    # superseded designs, kept as historical record.
    #
    # Regenerated deliberately: the ⚠️ SUPERSEDED banner was added to the
    # Markdown after these .docx files were made, so the shared copies carried
    # no warning at all — and they describe decisions the current specification
    # has since reversed, most importantly the frontend (ADR-0006). A document
    # that is wrong and does not say so is worse than one that is missing.
    "docs/design/PLATFORM_SPEC.md":     "docs/design/Nidan_Platform_Spec.docx",
    "docs/design/SYSTEM_DESIGN.md":     "docs/design/Nidan_System_Design.docx",
    # maps and status
    "docs/PROJECT_MAP.md":              "docs/Nidan_Project_Map.docx",
    "docs/build-log/STATUS.md":         "docs/build-log/Nidan_Build_Status.docx",
    "docs/build-log/README.md":         "docs/build-log/Nidan_Build_Log_Convention.docx",
    "docs/build-log/PHASE-1.md":        "docs/build-log/Nidan_Phase_1_Summary.docx",
}

# Every ADR becomes one document, because sixteen one-page files are useless as
# sixteen attachments.
ADR_DIR = "docs/spec/adr"
ADR_OUT = "docs/spec/Nidan_ADRs.docx"

MONO = "Consolas"


def build_documents() -> dict[str, str]:
    """The full map, including one entry per finished build log."""
    documents = dict(DOCUMENTS)
    for path in sorted((REPO / "docs" / "build-log").glob("T-*.md")):
        task = path.name.split("-", 2)[0] + "-" + path.name.split("-", 2)[1]
        documents[f"docs/build-log/{path.name}"] = \
            f"docs/build-log/Nidan_{task}_Build_Log.docx"
    return documents


# ── Markdown → Word ──────────────────────────────────────────────────

_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def add_runs(paragraph, text: str) -> None:
    """
    Inline formatting: **bold**, `code`, *italic*.

    Deliberately small. A full Markdown parser would handle nested emphasis and
    links, and nothing in these documents needs it — the cost of the missing
    cases is a stray asterisk, not a wrong document.
    """
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
            run.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows: list[str]) -> None:
    """A Markdown table, minus its separator row."""
    body = [r for r in rows if not re.fullmatch(r"\|[\s:|-]+\|", r.strip())]
    if not body:
        return
    cells = [split_row(r) for r in body]
    width = max(len(r) for r in cells)

    table = doc.add_table(rows=0, cols=width)
    table.style = "Light Grid Accent 1"
    for index, row in enumerate(cells):
        cursor = table.add_row().cells
        for column in range(width):
            text = row[column] if column < len(row) else ""
            paragraph = cursor[column].paragraphs[0]
            add_runs(paragraph, text)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True


def add_code(doc, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = MONO
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def convert(md_path: pathlib.Path, docx_path: pathlib.Path, *,
            title: str | None = None) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    heading = doc.add_heading(title or md_path.stem.replace("_", " "), level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    stamp = doc.add_paragraph()
    run = stamp.add_run(
        f"Nidan · generated from {md_path.relative_to(REPO).as_posix()} "
        f"on {date.today().isoformat()}")
    run.italic = True
    run.font.size = Pt(8.5)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code(doc, block)

        elif stripped.startswith("|") and stripped.endswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(lines[index])
                index += 1
            add_table(doc, rows)
            continue

        elif re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            doc.add_paragraph()

        elif stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            add_runs(doc.add_heading(level=min(level, 4)),
                     stripped.lstrip("#").strip())

        elif re.match(r"^\s*[-*+]\s+", line):
            add_runs(doc.add_paragraph(style="List Bullet"),
                     re.sub(r"^\s*[-*+]\s+", "", line))

        elif re.match(r"^\s*\d+[.)]\s+", line):
            add_runs(doc.add_paragraph(style="List Number"),
                     re.sub(r"^\s*\d+[.)]\s+", "", line))

        elif stripped.startswith(">"):
            quoted = stripped.lstrip(">").strip()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(24)
            if quoted.startswith("#"):
                # A heading inside a blockquote — which is how the SUPERSEDED
                # banners are written. Without this the "#" characters appear
                # literally, on the one line in those documents that most needs
                # to be read.
                add_runs(paragraph, quoted.lstrip("#").strip())
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(12)
            else:
                add_runs(paragraph, quoted)
                for run in paragraph.runs:
                    run.italic = True

        elif stripped:
            add_runs(doc.add_paragraph(), stripped)

        index += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def build_adrs(out: pathlib.Path) -> None:
    """All sixteen ADRs as one document."""
    combined: list[str] = ["# Architecture Decision Records", ""]
    for path in sorted((REPO / ADR_DIR).glob("0*.md")):
        combined.append(path.read_text(encoding="utf-8").strip())
        combined.extend(["", "---", ""])

    scratch = REPO / ADR_DIR / ".combined.md"
    scratch.write_text("\n".join(combined), encoding="utf-8")
    try:
        convert(scratch, out, title="Nidan — Architecture Decision Records")
    finally:
        scratch.unlink(missing_ok=True)


def main() -> None:
    wanted = {a.upper() for a in sys.argv[1:]}
    documents = build_documents()
    written = 0

    for source, target in sorted(documents.items()):
        stem = pathlib.Path(source).stem.upper()
        if wanted and stem not in wanted:
            continue
        md_path = REPO / source
        if not md_path.exists():
            print(f"  skipped (no source): {source}")
            continue
        convert(md_path, REPO / target)
        print(f"  {source}  ->  {target}")
        written += 1

    if not wanted or "ADR" in wanted or "ADRS" in wanted:
        build_adrs(REPO / ADR_OUT)
        print(f"  {ADR_DIR}/*.md  ->  {ADR_OUT}")
        written += 1

    print(f"\n{written} document(s) written.")


if __name__ == "__main__":
    main()
